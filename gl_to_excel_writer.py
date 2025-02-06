import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from database import Database
import json

class GlToExcelWriter:
    def __init__(self):
        # Load constants from JSON file
        with open('Configuration/constants.json') as f:
            constants = json.load(f)
        
        self.db_path = constants['gldbFilePath']
        self.excel_path = constants['excelWriter']['excelPath']
        self.sheet_name = constants['excelWriter']['sheetName']
        self.table_name = constants['excelWriter']['tableName']
        self.gl_items_table_name = constants['gldbGlItemsTableName']

    def write_gl_items_to_excel(self):
        """
        Reads the GL items from the database and adds them to a specified Excel sheet table.
        """
        gl_items = self._read_gl_items_from_db()
        self._write_and_add_table_to_excel(gl_items)

    def _read_gl_items_from_db(self):
        """
        Reads the GL items from the database and returns them as a DataFrame.
        """
        db = Database(self.db_path)
        query = f"SELECT * FROM {self.gl_items_table_name}"
        gl_items = pd.read_sql_query(query, db.connection)
        db.close()
        return gl_items

    def _write_and_add_table_to_excel(self, gl_items):
        """
        Writes the DataFrame data to the specified Excel sheet and adds a table.
        """
        with pd.ExcelWriter(self.excel_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            gl_items.to_excel(writer, sheet_name=self.sheet_name, index=False)

        book = load_workbook(self.excel_path)
        sheet = book[self.sheet_name]
        
        # Remove the existing table if it exists
        existing_tables = [tbl for tbl in sheet.tables.values() if tbl.displayName == self.table_name]
        for tbl in existing_tables:
            del sheet.tables[tbl.displayName]
        
        # Define the table range
        (max_row, max_col) = gl_items.shape
        table_range = f"A1:{chr(65 + max_col - 1)}{max_row + 1}"
        
        # Create a table
        table = Table(displayName=self.table_name, ref=table_range)
        
        # Add a default style with striped rows
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        
        # Add the table to the sheet
        sheet.add_table(table)
        
        # Set the column width for better visibility
        for i, col in enumerate(gl_items.columns):
            max_len = max(gl_items[col].astype(str).map(len).max(), len(col)) + 2
            sheet.column_dimensions[chr(65 + i)].width = max_len
        
        # Save the workbook
        book.save(self.excel_path)
